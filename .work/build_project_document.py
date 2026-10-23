from pathlib import Path
import sys
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'Görev ve Bağımlılık Yönetim Sistemi Proje Dokümanı v0.3.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
sec.footer_distance = Cm(.7)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Calibri')
    s.paragraph_format.space_after = Pt(6)
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.09
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
for name in ['List Bullet', 'List Number']:
    doc.styles[name].font.size = Pt(10.5)
    doc.styles[name].paragraph_format.space_after = Pt(4)
    doc.styles[name].paragraph_format.line_spacing = 1.05

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    return doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def bullets(items):
    for item in items:
        p(item, 'List Bullet')

def numbered(items):
    for i, item in enumerate(items, 1):
        p(f'{i}. {item}')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    for j, text in enumerate(headers):
        t.rows[0].cells[j].text = text
    for row in rows:
        cells = t.add_row().cells
        for j, text in enumerate(row):
            cells[j].text = text
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trpr.append(cant)
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            trpr.append(repeat)
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                elem = OxmlElement('w:' + edge)
                for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    elem.set(qn('w:' + k), v)
                borders.append(elem)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, val in [('top', '85'), ('bottom', '85'), ('left', '105'), ('right', '105')]:
                elem = OxmlElement('w:' + edge)
                elem.set(qn('w:w'), val)
                elem.set(qn('w:type'), 'dxa')
                margins.append(elem)
            tcpr.append(margins)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'E8EDF2' if i == 0 else ('F7F8FA' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.02
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    run.font.bold = i == 0
                    run.font.color.rgb = RGBColor(0,0,0)
    p('').paragraph_format.space_after = Pt(0)
    return t

def link(label, url):
    para = p('')
    rel = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyp = OxmlElement('w:hyperlink')
    hyp.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4D78'); rp.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); rp.append(size)
    run.append(rp)
    txt = OxmlElement('w:t'); txt.text = label; run.append(txt)
    hyp.append(run); para._p.append(hyp)
    para.paragraph_format.space_after = Pt(3)

foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = foot.add_run('Sürüm 0.3  |  '); r.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
doc.core_properties.title = 'Görev ve Bağımlılık Yönetim Sistemi'
doc.core_properties.subject = 'MVC katmanları ve modül sınırları ile AI ajan geliştirme şartnamesi'
doc.core_properties.author = ''

p('Görev ve Bağımlılık\nYönetim Sistemi', 'Title')
p('AI ajan geliştirme şartnamesi', 'Subtitle')
p('Sürüm 0.3  •  23 Ekim 2026')
h('Amaç ve kullanım')
p('Bu doküman, şirketlerin Microsoft 365 hesaplarıyla kullandığı görev ve bağımlılık yönetim sisteminin ilk sürümünü tanımlar. Yazılımcı veya AI ajan; kapsamı, iş kurallarını ve modül sınırlarını esas alarak geliştirme yapar. Yeni kullanıcı talepleri ilgili modül ve kabul senaryoları güncellenerek ele alınır.')
p('Sistem, işlerin hangi görevi beklediğini ve gecikmenin nereden kaynaklandığını gösterir. Proje planları gerçekleşen tarihlerle karşılaştırılır; atanan kullanıcıların görev performansı raporlanır. Bulgular görev, atama ve zaman kayıtlarına dayanır; bağımlılıktan bekleme ayrı gösterilir.')
h('İlk sürümün kapsamı')
bullets([
    'Şirket bazında veri ayrımı ve Microsoft 365 / Entra ID ile oturum açma.',
    'Proje oluşturma, yetkilendirme, görev atama ve planlanan / gerçekleşen tarihleri izleme.',
    'Öncül ve ardıl görev ilişkileri, döngü kontrolü ve öncüller bitince görevin hazır hâle gelmesi.',
    'Gecikme kaynağının ve bağımlılık nedeniyle bekleyen görevlerin gösterilmesi.',
    'Gecikme açıklaması takibi, site içi bildirim, Web Push ve Microsoft Graph üzerinden e-posta.',
    'Proje bazında planlanan / gerçekleşen tarihlerin, tamamlanmanın ve sapmaların raporlanması; atanan kullanıcıların görev performansı ve toplam puanı.',
    'iOS, Android, Windows ve macOS üzerinde mobil uyumlu web uygulaması ve ana ekrana kurulabilen PWA.'
])
h('Korunan teknoloji yönü')
p('İstemci: mobil uyumlu web ve PWA. Kimlik: Entra ID, MSAL, OAuth 2.0 / OIDC. Sunucu: Node.js ve TypeScript; API Gateway, mikroservisler ve arka plan işleyicileri. Veri: PostgreSQL ve RLS; önbellek / kuyruk: Redis. Bildirim: site içi, standart Web Push ve Microsoft Graph e-postası.')
p('Web arayüz frameworkü, sunucu MVC frameworkü, Kong / Envoy ve kuyruk uygulaması kodlamadan önce seçilir. İlk sürüm tarayıcı ve PWA üzerinden kullanılır. Ayrıntılar ve açık kararlar son sayfadadır.')

page('MVC katmanları ve modül sınırları')
p('MVC, her iş modülünün sorumluluklarını ayırmak için kullanılır. Mikroservis dağıtımı bu mantıksal yapının üzerindedir; her modül için ayrı servis, veritabanı veya kuyruk zorunluluğu oluşturulmaz. Kesin servis gruplaması kodlamadan önce kararlaştırılır.')
table(['Katman', 'Sorumluluk', 'Sınır'], [
    ('View', 'Mobil uyumlu web ekranlarını ve kullanıcı girdisini yönetir.', 'Sunucudaki iş ve yetki kurallarını değiştiremez.'),
    ('Controller', 'İsteği alır, biçimini doğrular, servisi çağırır.', 'SQL, bağımlılık veya gecikme algoritması içermez.'),
    ('Uygulama servisi', 'Kullanım senaryosu, yetki ve işlem akışını yürütür.', 'Diğer modüllere tanımlı arayüzlerle ulaşır.'),
    ('Model ve domain', 'Varlıkları ve iş kurallarını tanımlar.', 'HTTP, ekran ve e-posta sağlayıcısını bilmez.'),
    ('Veri erişimi', 'Repository üzerinden kalıcı veri okur / yazar.', 'İş kararı vermez; şirket sınırını korur.'),
    ('Altyapı', 'DB, MSAL / Graph, Redis ve zamanlayıcı bağlantıları.', 'İş kurallarının sahibi değildir.')
], [3.0, 7.0, 7.4])
h('İş modülleri')
table(['Modül', 'Sahip olduğu alan'], [
    ('Kimlik ve şirket', 'Oturum, kullanıcı eşleştirme ve güvenilir şirket bağlamı.'),
    ('Proje ve yetki', 'Proje, üyelik, roller, görünürlük talebi ve erişim kararı.'),
    ('Görev', 'Görev, atama, tarihler ve bütün çalışma durumu geçişleri.'),
    ('Bağımlılık', 'Öncül–ardıl ilişkileri, döngü kontrolü ve hazır olma hesabı.'),
    ('Gecikme ve açıklama', 'Gecikme teşhisi, gerekçe ve açıklama takibi.'),
    ('Bildirim', 'Site içi / e-posta / Web Push, abonelikler ve gönderim tekrarları.'),
    ('Raporlama', 'Proje plan / gerçekleşen karşılaştırmaları, kullanıcı göstergeleri ve puanı.')
], [4.6, 12.8])
h('Değişikliklerin sınırlandırılması')
p('Her modül kendi kayıtlarını yazar. Başka modülün tablosuna, repository uygulamasına veya iç modeline doğrudan erişilmez. Modüller servis sözleşmeleri ve olay mesajları üzerinden konuşur. Görev durumunu yalnız Görev modülü değiştirir; Bağımlılık ve Gecikme modülleri bu modülün servislerini çağırır.')
p('Paylaşılan alan; kimlik türleri, şirket bağlamı, ortak hata biçimi ve sözleşmelerle sınırlı tutulur. İş kuralları ortak yardımcı dosyalara taşınmaz. İstemci ve sunucuda görünürlük aynı sonucu üretir; son karar sunucudadır.')
p('Raporlama modülü, veri sahibi modüllerin salt okunur sözleşmelerini kullanır. Görev durumu, tarihler, atamalar veya gecikme gerekçelerini değiştiremez. Rapor görünürlüğünü Proje ve yetki modülü belirler; hesaplar ekran veya controller içine yazılmaz.')

page('Yetkiler ve görev yaşam döngüsü')
h('Erişim modeli')
table(['Rol veya hak', 'Görünürlük', 'Değişiklik yetkisi'], [
    ('Proje yöneticisi', 'Yönettiği projenin tamamı.', 'Üyelik, roller ve proje / görev yönetimi.'),
    ('Editör', 'Yetkili olduğu projenin tamamı.', 'Proje, görev ve bağımlılık düzenleme; üyelik yönetimi ayrı yetkidir.'),
    ('İzleyici', 'Erişim verilen projenin tamamı.', 'Okuma hakkı; düzenleme hakkı vermez.'),
    ('Görev sorumlusu', 'Kendi görevleri ve gerekli bağımlılık özeti.', 'Kendi görevinin çalışma ilerlemesi ve açıklaması.'),
    ('Erişimsiz kullanıcı', 'Yetki verilmemiş proje ve görevleri göremez.', 'İşlem yapamaz.')
], [3.6, 6.8, 7.0])
p('Bir kişide izleyici hakkı ile görev sorumluluğu birlikte bulunabilir. Kullanıcı tüm projeyi görmek için talep açar; proje yöneticisi onayladığında izleyici hakkı kazanır ve kendi görevlerini güncelleme hakkını korur. Bekleyen veya reddedilen talep erişimi genişletmez. Talep düzenleme ya da proje yönetimi yetkisi vermez.')
p('Sınırlı bağımlılık özeti, işi engelleyen ilişkinin ve tamamlanma durumunun anlaşılmasını sağlar. Başka görevlerin tüm ayrıntılarını veya bütün bağımlılık zincirini açmaz. Başlık ve sorumlu gibi gösterilecek alanlar K06 ile kesinleşir. Aynı erişim kuralı liste, grafik, arama, bağlantı ve bildirimlerde uygulanır.')
p('Proje raporu mevcut proje görünürlüğüne tabidir. Performans ve puan için ayrı denetim yapılır: yönetici yalnız yönettiği projelerdeki kullanıcıları, kullanıcı kendisini görür. Editörlük / izleyicilik başka kişilerin raporunu açmaz. Kullanıcı geçmiş sorumluluğuna ait rapor verisini de görür; bu hak devredilmiş görevin güncel ayrıntılarını açmaz. API, toplam ve önbellek aynı sınırı korur.')
h('Çalışma durumları')
table(['Kod', 'Anlam', 'Geçiş koşulu'], [
    ('BLOCKED', 'Öncül görevler nedeniyle bekliyor.', 'Tamamlanmamış en az bir öncül vardır.'),
    ('READY', 'Başlamaya hazır.', 'Tamamlanmamış öncül yoktur.'),
    ('IN_PROGRESS', 'Sorumlu görevi yürütüyor.', 'Yetkili sorumlu başlatır; engel bulunamaz.'),
    ('COMPLETED', 'Görev tamamlandı.', 'Yetkili işlemle tamamlanma zamanı kaydedilir.')
], [3.5, 6.1, 7.8])
p('Yeni görev, öncüllerine göre BLOCKED veya READY oluşturulur. Temel akış READY → IN_PROGRESS → COMPLETED biçimindedir. Bir öncülün bitmesi, diğer öncüller bitmediyse görevi hazır yapmaz. Son tarihin geçmesi çalışma durumunu değiştirmez.')
p('Açıklama girmek görevi başlatmaz, engeli kaldırmaz ve tamamlanma zamanını değiştirmez. Açıklama beklerken tamamlama, tamamlanmış görevi yeniden açma ve başlamış göreve yeni bağımlılık ekleme davranışları K07 ile kararlaştırılır; ajan kendiliğinden kural eklemez.')

page('Bağımlılık gecikme ve bildirim akışları')
h('Bağımlılık kuralları')
bullets([
    'İlk sürümün ilişkisi Finish to Start biçimindedir: öncül tamamlanmadan ardıl başlayamaz.',
    'Görev kendisine bağlanamaz. A → B → C varken C → A eklenemez. Döngü kontrolü ve kayıt, eşzamanlı isteklerde de tutarlı kalır.',
    'Şirketler arası bağ kurulamaz. Projeler arası bağımlılık K05 kararı verilmeden uygulanmaz.',
    'Görev tamamlanınca doğrudan ardıllar değerlendirilir. Tüm öncülleri tamamlanan BLOCKED görev READY olur; diğer çalışma durumları körlemesine değiştirilmez.'
])
h('Gecikme ve açıklama takibi')
p('Günlük gecikme takibi, açıklama talebi ve bildirimler güncel plan bitişine göre çalışır. Performans ve ilk plan sapması ilk kaydedilen bitişi esas alır. Tarih revizyonu günlük takibi değiştirebilir; geçmiş performans sapmasını ve açıklama kayıtlarını silmez.')
numbered([
    'Tamamlanmamış görev güncel plan bitişini geçince operasyonel gecikme kaydı açılır. Tamamlama sırasında da güncel bitiş hedefi kontrol edilir; gecikme kaydı eksikse oluşturulur. Worker taramasından önce geç tamamlanan görev kaçırılmaz. İlk plana göre sapma raporda ayrıca hesaplanır.',
    'Görev bağımlılık nedeniyle bekliyorsa sistem neden olan görevleri ve tespit zamanını kaydeder. Bu beklemenin gerekçesi otomatik karşılanır; aynı bekleme için kullanıcıdan tekrar mazeret istenmez.',
    'Bağımlılıkla açıklanamayan gecikmede sorumludan açıklama beklenir. Açıklama girilene kadar belirlenen sıklıkta hatırlatma yapılır.',
    'Sorumlu açıklamayı girince metin, kullanıcı ve zaman saklanır; ilgili proje yöneticilerine bildirim oluşturulur. O açıklama talebinin hatırlatmaları durur.',
    'Engel kalkınca çalışma durumu yeniden değerlendirilir; önceki sistem gerekçesi silinmez. Sonraki gecikme için yeni açıklama talebinin ne zaman açılacağı K04 kararıdır.'
])
p('Açıklama, görev durumundan ayrı saklanır: REQUIRED açıklama beklendiğini, RECORDED gerekçenin kaydedildiğini belirtir. Kaynak SYSTEM veya USER olur. Örneğin bir görev aynı anda BLOCKED ve RECORDED / SYSTEM olabilir. EXPLANATION_REQUIRED çalışma durumu olarak kullanılmaz.')
h('Gecikme kaynağının gösterimi')
p('A geciktiği için B başlayamıyorsa sistem A’yı gecikme kaynağı, B’yi bağımlılıktan etkilenen görev olarak gösterir. Birden fazla kaynak listelenebilir. Önceden geç bitmiş herhangi bir görev, mevcut gecikmenin tek nedeni sayılmaz. Bağlarda engelleyici görev ile bekleme başlangıç / bitiş zamanı tutulur; gecikme kaydı bu kanıtı korur. Tarihlerle neden ilişkisi kurulamıyorsa sonuç kesinleştirilmez.')
h('Bildirim akışı')
p('Görev hazır olduğunda sorumluya; güncel son tarihten 24 saat önce bir kez; açıklama beklenirken günde bir; açıklama girilince yöneticilere bildirim oluşturulur. Sistem gerekçesiyle karşılanan bekleme için manuel açıklama hatırlatması gönderilmez. Plan revizyonunda bekleyen eski tarih bildirimleri yeniden planlanır; gönderim geçmişi korunur. Tarihler kendiliğinden ötelenmez. Saat dilimi / iş günü hesabı K03 kararıdır.')
p('Bildirim, iş değişikliği kaydedildikten sonra site içi, e-posta ve izinli Web Push kanallarında işlenir. Kanal hatası görev işlemini geri almaz. Gönderim öncesi güncel alıcı yetkisi kontrol edilir. Tekrar deneme / mükerrerlik Bildirim modülündedir. Graph veya push servisinin isteği kabul etmesi, kullanıcıya teslim garantisi sayılmaz.')

page('Temel veri modeli ve servis sözleşmeleri')
p('Aşağıdaki model asgari ilişkileri tanımlar. Alan türleri, indeksler ve migration dosyaları seçilen sunucu frameworküyle tamamlanır. Kalıcı iş kayıtları şirket kimliği taşır; ilişkili kayıtların aynı şirkete ait olması doğrulanır.')
table(['Varlık', 'Asgari alanlar ve ilişki', 'Sahibi'], [
    ('Tenant', 'id, ad, entra_tenant_id', 'Kimlik ve şirket'),
    ('User', 'id, tenant_id, entra_oid, ad, e-posta', 'Kimlik ve şirket'),
    ('Project', 'id, tenant_id, başlık, açıklama, ilk / güncel plan başlangıç ve bitişi; gerçekleşen tarihler görevlerden türetilir', 'Proje ve yetki'),
    ('ProjectPlanChange', 'tenant_id, project_id, önceki / yeni plan tarihleri, değiştiren kullanıcı / zaman', 'Proje ve yetki'),
    ('ProjectMember', 'tenant_id, project_id, user_id, rol', 'Proje ve yetki'),
    ('VisibilityRequest', 'tenant_id, project_id, user_id, talep zamanı, durum, karar veren yönetici / zaman', 'Proje ve yetki'),
    ('Task', 'id, tenant_id, project_id, assignee_id, başlık, work_status, ilk / güncel plan ve gerçekleşen tarihler, tamamlayan kullanıcı', 'Görev'),
    ('TaskAssignment', 'tenant_id, task_id, user_id, atama başlangıç / bitişi, işlemi yapan kullanıcı', 'Görev'),
    ('TaskPlanChange', 'tenant_id, task_id, önceki / yeni plan tarihleri, değiştiren kullanıcı / zaman', 'Görev'),
    ('TaskDependency', 'tenant_id, predecessor_id, successor_id, ilişki türü, bekleme başlangıç / bitiş zamanı', 'Bağımlılık'),
    ('DelayRecord', 'tenant_id, task_id, tespit zamanı, açıklama durumu / kaynağı, gerekçe, engel zaman aralıkları, neden sınıfı / kanıtı, bildiren kullanıcı / zaman', 'Gecikme ve açıklama'),
    ('Notification', 'tenant_id, alıcı, olay kimliği, kanal, push aboneliği, gönderim durumu, deneme sayısı / zaman', 'Bildirim'),
    ('PushSubscription', 'id, tenant_id, user_id, endpoint, p256dh, auth, oluşturulma / iptal zamanı', 'Bildirim')
], [3.6, 10.4, 3.4])
p('User(tenant_id, entra_oid), üyelik(tenant_id, project_id, user_id) ve bağ(tenant_id, predecessor_id, successor_id) benzersizdir. İlişkiler aynı şirkette kalır. Atama, plan değişikliği, bağ kaldırma veya açıklama girişi geçmiş rapor kanıtını silmez. Aynı görevin atama aralıkları çakışamaz; ilk plan ilk kayıt anında ayrıca korunur.')
h('Modüller arası sözleşmeler')
bullets([
    'Kimlik ve şirket → güvenilir kullanıcı / şirket bağlamı. İstemciden gelen tenant_id tek başına yetki sayılmaz.',
    'Proje ve yetki → görünür projeler / görevler ve işlem izni. Görev sorumlusu için proje üyeliği zorunlu varsayılmaz.',
    'Bağımlılık → bağ ekleme / kaldırma ve engelleyici öncülleri bulma.',
    'Görev → oluşturma, atama, başlatma, tamamlama ve bağımlılığa göre hazır olma durumunu güncelleme.',
    'Gecikme → gecikmeyi değerlendirme, sistem gerekçesi kaydetme ve kullanıcı açıklaması alma.',
    'Bildirim → olaydan bildirim oluşturma ve gönderimi işleme; şirket / alıcı sınırını koruma.',
    'Raporlama → proje karşılaştırması, kullanıcı göstergeleri ve puan; plan, veri kesimi, kapsam ve formül sürümüyle birlikte. Veri sahibi modüller güncel ve geçmiş sorumluluk için yetkili rapor verisi sunar.'
])
p('Tutarlılık gerektiren değişiklikler birlikte yürütülür. Ayrı mikroservislerde kontrolsüz çift yazım yapılmaz; olay kaydı ve tekrar işleme servis dağılımıyla belirlenir. PostgreSQL RLS etkinleştirilir; uygulama hesabı RLS’yi atlayamaz. Önbellek de şirket ve erişim kapsamını ayırır.')

page('Web ve PWA kullanımı ile bildirimler')
h('İstemci kapsamı ve mobil kullanım')
p('İlk sürüm, telefon ve masaüstüne uyumlu tek bir web uygulamasıdır. PWA kurulumu ve Web Push destekleyen tarayıcılarda bu özellikler açılır; temel proje, görev ve rapor ekranları tarayıcıdan kullanılabilir. PWA seçimi çevrimdışı görev düzenleme veya senkronizasyon kapsamı oluşturmaz.')
p('Uygulama HTTPS üzerinden sunulur. Manifest, uygulama kimliği, ad, ikonlar ve bağımsız pencere görünümünü tanımlar. Service worker Web Push alımını ve bildirim gösterimini yürütür. iOS 16.4 ve üzerindeki kullanıcı uygulamayı ana ekrana ekler, oradan açar ve “Bildirimleri aç” gibi bir düğmeye dokunarak izin verir. Standart Web Push için Apple Developer Program üyeliği gerekmez.')
p('Android’de Chrome gibi Web Push destekleyen bir tarayıcı üzerinden bildirim alınır. HTTPS sitesine verilen bildirim izni ile Android’in tarayıcı / site bildirim ayarları açık olmalıdır. Android Chrome’da ana ekrana PWA kurulumu ön koşul değildir. Kullanıcı izin verdikten sonra sayfa açık değilken de bildirim alınabilir; ağ, cihaz ve sistem ayarları teslim zamanını etkileyebilir.')
p('Web Push, destek ve kullanıcı iznine bağlıdır. İzin verilmemesi temel kullanımı engellemez; site içi bildirim ve e-posta kanalları devam eder. iPhone’da ana ekrana ekleme adımı arayüzde açıklanır. Push desteklenmiyorsa veya kapatılmışsa kullanıcı mevcut kanallardan bildirimlerini takip eder.')
h('Gönderim ve abonelik yönetimi')
p('Sunucudaki worker / Redis; 24 saatlik hatırlatmayı, günlük açıklama hatırlatmasını ve tekrar denemeleri zamanlar. Tarayıcıdaki service worker gelen push’u gösterir. Zamanlama açık bir sekmeye veya tarayıcıda sürekli çalışan bir sayaca bağlanmaz. Standart Web Push aboneliği ve VAPID anahtarları kullanılır; özel anahtar sunucuda tutulur.')
p('Bir kullanıcı birden fazla cihaz / tarayıcı aboneliğine sahip olabilir. Abonelik şirket ve kullanıcıya sunucudaki oturumdan bağlanır. Çıkışta veya hesap değişiminde önceki kullanıcı bağlantısı devre dışı bırakılır. Geçersiz / iptal edilmiş aboneliğe tekrar gönderim durdurulur. Mükerrerlik olay + kanal + alıcı / abonelik bazında kontrol edilir.')
p('Gönderimden önce alıcının güncel erişimi doğrulanır. Kilit ekranı mesajı sınırlı tutulur; ayrıntılar görev API’sinden alınır. Bildirime tıklama ilgili görev yolunu açar; gerekirse giriş yapılır ve güncel yetki yeniden kontrol edilir. Bildirim bağlantısı erişim hakkı sağlamaz. Kişisel görev ve rapor yanıtları service worker önbelleğine otomatik alınmaz.')
h('Web ve PWA kabul senaryoları')
table(['Kod', 'Senaryo ve beklenen sonuç'], [
    ('WEB01', 'iOS 16.4+ cihazda ana ekrana kurulum ve düğmeyle izin sonrası Web Push görünür. İzin reddinde temel ekranlar, site içi bildirim ve e-posta çalışır.'),
    ('WEB02', 'Birden fazla abonelik doğru kullanıcı / şirkete bağlıdır. Çıkış veya hesap değişimi eski bağlantıyı kapatır; geçersiz abonelik pasifleştirilir.'),
    ('WEB03', 'Görev devri / yetki iptali sonrası kuyruktaki gönderim güncel yetkiyi kontrol eder. Tıklama ile yetkisiz görev ayrıntısı veya eski kullanıcı önbelleği açılmaz.'),
    ('WEB04', 'Worker tekrarında aynı olay abonelik bazında kontrolsüz çoğalmaz. Kapalı sayfada da sunucu zamanlaması yürür; push hatası görevi veya diğer kanalları bozmaz.'),
    ('WEB05', 'Gerçek Android cihazda PWA kurulmadan Chrome’dan izin verilir; sayfa kapatılınca test bildirimi doğrulanır. Tıklamada oturum / güncel görev yetkisi kontrol edilir. Site / sistem izni kapalıysa durum doğru gösterilir; temel kullanım sürer.')
], [1.9, 15.5])

page('Proje raporları ve kullanıcı göstergeleri')
h('Proje planı ve gerçekleşen sonuçlar')
p('Projenin planlanan başlangıç ve bitişi ayrıca girilir. İlk kaydedilen proje ve görev planları korunur; sonradan yapılan değişiklikler güncel planda ve değişiklik geçmişinde tutulur. Rapor, ilk planı, güncel planı ve gerçekleşeni yan yana gösterir. Tarih revizyonu ilk plana göre sapmayı ortadan kaldırmaz.')
p('Projenin gerçekleşen başlangıcı ilk başlayan görevin başlangıcıdır. Gerçekleşen bitiş, bütün görevler tamamlandığında son tamamlanma zamanıdır. Görev başlamadıysa başlangıç, proje tamamlanmadıysa gerçekleşen bitiş boş kalır; tahmin veya rapor tarihiyle doldurulmaz. Görevsiz projede gerçekleşen tarih ve tamamlama oranı hesaplanamaz.')
p('Proje raporu başlangıç / bitiş tarihlerini, görev adetlerini, açık ve gecikmiş işleri, adet bazında tamamlanma oranını ve görev bazındaki tarih sapmalarını içerir. Tamamlanma oranı tamamlanan görev / rapora dahil tüm görevlerdir; iş gücü veya emek yüzdesi anlamına gelmez. Görev ekleme / çıkarma ve yeniden açmanın geçmiş rapora etkisi K09 ile tamamlanır.')
h('Dönem ve veri kesimi')
p('Rapor şirket, proje, kullanıcı ve dönemle filtrelenir. Dönem başlangıcı dahil, bitişi hariçtir. Her sonuç veri kesim zamanını gösterir. Hedef görevler ilk plan bitişine, tamamlanan işler gerçekleşen bitişine, atamalar atama başlangıcına göre dönemlenir. Açık işler rapor üretim anının fotoğrafıdır; bu üç küme aynı toplam gibi sunulmaz. Saat dilimi K03 ile kesinleştirilir.')
table(['Gösterge', 'Hesap ve ilişkilendirme'], [
    ('Atanan görev', 'Dönemde atanan benzersiz görev sayısı; atama geçmişinden kullanıcı bazında.'),
    ('Tamamlanan iş', 'Dönemde biten görev sayısı; bitiş anındaki görev sorumlusuna yazılır. İşlemi yapan yönetici sahip sayılmaz.'),
    ('Zamanında tamamlama', 'İlk plan vadesi dönemde olan ve veri kesiminde vadesi gelmiş görevlerde, zamanında biten / tüm hedefler. Vade anındaki sorumlu esas alınır.'),
    ('Açık ve gecikmiş iş', 'Kesim anındaki açık işler ve ilk plan vadesini aşanlar ayrı gösterilir. Açık iş yükü güncel sorumluya aittir.'),
    ('Tarih sapması', 'Tamamlanan işte gerçek bitiş − ilk plan bitişi; açık gecikmede kesim zamanı − ilk plan bitişi. Erken / geç sonuç ve güncel plan farkı ayrıca gösterilir.'),
    ('Bağımlılık beklemesi', 'Kanıtlanan bekleme aralıklarının dönemle kesişen birleşim süresi. Çakışan engeller iki kez sayılmaz; görev ve kullanıcının atama aralığı ayrımı korunur.')
], [4.1, 13.3])
p('Geçmişte teslimi gereken fakat bu dönemde biten görev, bu dönemin tamamlanan iş sayısına girer; bu dönemin teslim hedeflerine taşınmaz. Ham göstergeler tüm yetkili görevleri içerir. Puan için uygulanan bağımlılık istisnası, ham gecikme ve proje sapması kayıtlarını silmez.')

page('Kullanıcı performans puanı ve rapor kabulü')
h('Puanın hesaplanması')
p('Her kullanıcı için H, ilk plan bitişi seçilen dönemde bulunan ve kesim anında vadesi gelmiş görevlerdir. Her görev ilk plan vadesi anındaki sorumluya bağlanır. Henüz vadesi gelmeyen görevler puan hesabına girmez. Bitiş anındaki sorumluya yazılan tamamlanan iş hacmi bu kümenin yerine kullanılmaz.')
p('H içinden, gecikmesinin bütünüyle bağımlılıktan kaynaklandığı kanıtlanan görevler çıkarılarak S kümesi oluşturulur. Çıkarılan görev adedi ve nedeni ayrıca gösterilir. Yalnız BLOCKED geçmişi bulunması yeterli değildir; kısmi veya belirsiz nedenler otomatik muafiyet almaz. Tam bağımlılık kanıtının sınıflandırma ölçütü K10 ile kesinleştirilir.')
p('Z = S içinde ilk plan vadesine kadar tamamlanan görev sayısı / S görev sayısı. T = S içinde veri kesim zamanına kadar tamamlanan görev sayısı / S görev sayısı. İki oran 0 ile 1 arasındadır ve aynı kullanıcı / proje kapsamıyla hesaplanır.')
p('Toplam performans puanı = 100 × (0,70 × Z + 0,30 × T).')
p('S boşsa puan 0 veya 100 verilmez; “hesaplanamıyor” gösterilir. Örneğin S içinde 10 görevden 8’i zamanında, toplam 9’u tamamlanmışsa puan 83 olur. Kalan görevin kesimden sonra tamamlanması, sonraki rapor üretiminde T değerini değiştirebilir; hesap tarihi sonuçta görünür.')
p('Birden fazla projedeki puan, proje puanlarının basit ortalamasıyla hesaplanmaz; yetkili S görevlerinin ortak pay / paydasından hesaplanır. Yönetici yönetim kapsamını, kullanıcı ilgili göstergeye göre güncel veya geçmiş sorumluluğunu görür. Devir önceki sahibin puan kaydını silmez; rapor verisi güncel atamayla sınırlandırılmaz. Görev sayısı, kapsam ve hariç tutulanlar gösterilir. Bekleme gecikmeden çıkarılarak kişisel net gecikme üretilmez.')
h('Raporlamanın kabul senaryoları')
table(['Kod', 'Senaryo ve beklenen sonuç'], [
    ('RAP01', 'Plan tarihi değişince ilk plan korunur. İlk plan, güncel plan ve gerçekleşen karşılaştırmaları farklı sonuçlarını doğru gösterir.'),
    ('RAP02', 'Devirde bitişteki sorumlu tamamlanan işi, ilk vade anındaki sorumlu hedef başarısını alır. Tamamlama düğmesine basan yönetici performans sahibi olmaz.'),
    ('RAP03', 'Önceki dönemin hedefi bu dönemde bitince yalnız bu dönemin tamamlanan iş sayısına girer. Gelecek vadeler puanı düşürmez; boş payda hesaplanamıyor verir.'),
    ('RAP04', 'Bütünüyle bağımlılıktan geciken görev puanın pay ve paydasından çıkar; adedi görünür. Kısmi / belirsiz neden otomatik çıkarılmaz; çakışan beklemeler tek sayılır.'),
    ('RAP05', '8 zamanında ve 9 tamamlanmış / 10 uygun görev için puan 83’tür. Projeler arası puan uygun görevler üzerinden birleşir; basit proje ortalaması alınmaz.'),
    ('RAP06', 'İlk görev başlamadan proje gerçek başlangıcı, tüm görevler bitmeden gerçek bitişi oluşmaz. Bağımlılık ve açıklama işlemleri geçmiş sapmayı silmez.'),
    ('RAP07', 'Yönetici yönetim kapsamını, kullanıcı kendi güncel / geçmiş raporunu görür. Devir eski puanı silmez ve görevin güncel ayrıntısını açmaz. API / toplam / önbellek bu sınırı korur.')
], [1.9, 15.5])

page('AI ajan geliştirme ve kabul kuralları')
h('Her talepte izlenecek yol')
numbered([
    'Talebi mevcut kapsamla eşleştir. Etkilenen modülleri, servis sözleşmelerini ve kabul senaryolarını belirle; rapor hesabı ve veri geçmişine etkisini kontrol et.',
    'Belirsiz iş kuralı, yeni yetki veya kapsam değişikliği varsa ürün sahibine sor. Yalnız bu karara bağlı geliştirmeyi beklet; bağımsız işleri sürdür.',
    'İş mantığını domain ve uygulama servisinde uygula. Controller, ekran veya başka modülün repository’si içine taşıma.',
    'Veri değişiyorsa migration ve mevcut kayıtların korunmasını ele al. Ortak sözleşme değişiyorsa kullanan modülleri aynı değişiklikte güncelle.',
    'Değişen iş kuralını, modüller arası etkiyi ve erişim sınırlarını anlamlı testlerle doğrula. İlgisiz modüllerde gereksiz düzenleme yapma.',
    'Dokümandaki ilgili kuralı ve kabul senaryosunu güncelle. Sonuçta değişen davranışı, etkilenen modülleri, yapılan doğrulamayı ve kalan kararları bildir.'
])
h('İlk sürümün kabul senaryoları')
table(['Kod', 'Senaryo ve beklenen sonuç'], [
    ('KAB01', 'A şirketinin kullanıcısı B şirketinin proje veya görev kimliğini bilse bile okuyamaz / değiştiremez.'),
    ('KAB02', 'Görev sorumlusu kendi görevini ve sınırlı bağımlılık özetini görür; başka görevlerin ayrıntılarına ve üye listesine erişemez.'),
    ('KAB03', 'Görünürlük talebi onay öncesi erişimi artırmaz. Yönetici onayından sonra sorumlu tüm projeyi okur; yalnız kendi görev ilerlemesini ve açıklamasını günceller.'),
    ('KAB04', 'A → B → C varken C → A ve A → A reddedilir. Eşzamanlı bağımlılık istekleri döngü oluşturamaz.'),
    ('KAB05', 'C, A ve B’yi bekliyorsa yalnız A’nın bitmesi yeterli olmaz. B de bitince C bir kez READY olur.'),
    ('KAB06', 'Geciken BLOCKED görevde sistem gerekçesi oluşur; gerekçe görevi IN_PROGRESS yapmaz ve başka görevin gizli bilgisini açmaz.'),
    ('KAB07', 'Manuel açıklama gereken görevde sorumlu açıklama girince kayıt korunur, yöneticiler bilgilendirilir ve o talebin hatırlatmaları durur.'),
    ('KAB08', 'Geçmişte geç biten bir öncül, hedef görevin bugünkü gecikmesini otomatik açıklamaz. Teşhis ilgili engel kayıtlarını esas alır.'),
    ('KAB09', 'Aynı olay veya worker işi tekrar işlendiğinde görev durumu bozulmaz; aynı bildirim kontrolsüz çoğalmaz. E-posta hatası tamamlamayı geri almaz.'),
    ('KAB10', 'Açıklama çalışma durumunu değiştirmez. Worker çalışmadan güncel planı aşarak tamamlanan görevde de gecikme kaydı oluşur; ilk plan sapması raporda ayrıca korunur.'),
    ('KAB11', 'Güncel son tarihten 24 saat önce bir bildirim, açık açıklama talebinde günde bir hatırlatma oluşur. Plan revizyonunda eski tarihli işler yenilenir; otomatik tarih öteleme yapılmaz.')
], [1.9, 15.5])
p('RAP01–RAP07 raporlama, WEB01–WEB05 web / PWA bölümündedir. Açık karar kesinleşince ilgili test tamamlanır. Desteklenen tarayıcılar, M365 oturumu, Web Push ve Graph gönderimi gerçek ortamla doğrulanır.')

page('Uygulama öncesi kararlar ve teknik başvurular')
p('Aşağıdaki kararlar ürün sahibiyle kesinleştirilir. Bir kararın açık olması kapsamın genişletilmesi için izin oluşturmaz. Karar kaydı güncellendikten sonra ilgili iş kuralı, veri modeli ve kabul senaryosu birlikte revize edilir.')
table(['Karar', 'Kesinleştirilecek konu'], [
    ('K02', 'Web arayüz frameworkü ve M365 web oturum akışı; Node.js / TypeScript MVC frameworkü; Kong / Envoy, Redis kuyruk uygulaması ve servis gruplaması.'),
    ('K03', 'Şirket / proje saat dilimi, takvim veya iş günü hesabı ve günlük hatırlatmanın gönderim saati. 24 saat / günde bir ve otomatik ötelememe kuralları onaylıdır.'),
    ('K04', 'Engel kalktıktan sonra yeni manuel açıklama talebinin zamanı ve güncel plan revizyonunun açık açıklama talebine etkisi. Önceki gerekçe korunacak.'),
    ('K05', 'Bağımlılıklar yalnız aynı projede mi, aynı şirketin farklı projeleri arasında da kurulabilecek mi?'),
    ('K06', 'Sınırlı bağımlılık özetinde görev adı / kodu, sorumlu ve tarih alanlarından hangileri gösterilecek?'),
    ('K07', 'Açıklama beklerken tamamlama, görevi yeniden açma ve başlamış göreve bağımlılık ekleme kuralları.'),
    ('K08', 'Şirketin ilk kaydı ve yönetici atanması; Graph gönderici hesabı ve izinleri; desteklenecek tarayıcılar ve minimum sürümleri. iOS Web Push alt sınırı 16.4’tür.'),
    ('K09', 'Sonradan eklenen / çıkarılan veya yeniden açılan görevlerin geçmiş rapor hedeflerine ve proje gerçekleşen bitişine etkisi.'),
    ('K10', 'Gecikmenin bütünüyle bağımlılıktan kaynaklandığını kanıtlayan kesin ölçüt ve kayıttan belirlenemeyen nedenlerin işaretlenmesi. Puan formülü ve tam bağımlılık istisnası onaylıdır.')
], [1.9, 15.5])
h('Teknik uygulama notları')
p('Web oturumu Entra ID / MSAL ile uygulanır; akış seçilen web frameworküyle netleştirilir. SPA seçilirse authorization code + PKCE kullanılır. Graph Mail.Send izin modeli ve şirket yöneticisi onayı uygulama öncesinde seçilir. RLS, tenant filtrelerinin yerine tek başına bırakılmaz.')
p('Otomatik hatırlatmalar, gecikme kontrolleri ve başarısız gönderimlerin tekrar denenmesi Node.js / TypeScript arka plan işleyicisi ve Redis kuyruğuyla yürütülür. Bu işler proje görevlerinin tarih ve durum bilgilerini kullanır.')
h('Resmî teknik başvurular')
link('Apple web uygulamalarında Web Push', 'https://developer.apple.com/documentation/usernotifications/sending-web-push-notifications-in-web-apps-and-browsers')
link('WebKit iOS ve iPadOS ana ekran uygulamalarında Web Push', 'https://webkit.org/blog/13878/web-push-for-web-apps-on-ios-and-ipados/')
link('Service Worker ve HTTPS gereksinimi', 'https://developer.mozilla.org/en-US/docs/Web/API/Service_Worker_API')
link('Web Push aboneliği ve anahtarları', 'https://developer.mozilla.org/en-US/docs/Web/API/PushSubscription')
link('Google Chrome Android site bildirimleri ve izin ayarları', 'https://support.google.com/chrome/answer/3220216?co=GENIE.Platform%3DAndroid&hl=tr')
link('MDN PWA arka plan çalışması ve sayfa kapalıyken Web Push', 'https://developer.mozilla.org/en-US/docs/Web/Progressive_web_apps/Guides/Offline_and_background_operation')
link('Microsoft MSAL kimlik doğrulama akışları', 'https://learn.microsoft.com/en-us/entra/identity-platform/msal-authentication-flows')
link('PostgreSQL satır düzeyinde güvenlik', 'https://www.postgresql.org/docs/current/ddl-rowsecurity.html')
link('Microsoft Graph sendMail', 'https://learn.microsoft.com/en-us/graph/api/user-sendmail?view=graph-rest-1.0')

doc.save(OUT)
text = '\n'.join(p.text for p in doc.paragraphs)
tabletext = '\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
(ROOT / '.work' / 'document_text.txt').write_text(text + '\n' + tabletext, encoding='utf-8')
print(json.dumps({'path':str(OUT), 'paragraphs':len(doc.paragraphs), 'tables':len(doc.tables), 'words':len((text + tabletext).split())}, ensure_ascii=False))
